"""Read a .docx and turn it into simple, labelled items.

No AI here: just rules based on style names, bold/italic and a few patterns.
Every paragraph becomes an "item" with a role (bullet, entry_title, ...) and
named fields (text, label, value, part1, part2 ...).
"""
import re
from docx import Document
from docx.oxml.ns import qn

from .headings import canonical

PART_SPLIT = re.compile(r"\s+\|\s+")
CONTACT_SPLIT = re.compile(r"\s*\|\s*")
META_RE = re.compile(
    r"^(tech stack|technologies|technology|tools|duration|dates|date|"
    r"repository|repo|github|link)\s*:\s*(.*)$",
    re.I,
)
META_ROLE = {
    "tech stack": "tech", "technologies": "tech", "technology": "tech",
    "tools": "tech", "duration": "duration", "dates": "duration",
    "date": "duration", "repository": "repo", "repo": "repo",
    "github": "repo", "link": "repo",
}
BULLET_CHARS = "\u2022\u00b7\u25cf\u25aa\u2013\u2014*-"


# ---------- low-level reading ----------

def _flag(rpr, tag):
    if rpr is None:
        return None
    el = rpr.find(qn(tag))
    if el is None:
        return None
    return el.get(qn("w:val")) not in ("0", "false", "off")


def _style_flag(style, attr):
    while style is not None:
        value = getattr(style.font, attr)
        if value is not None:
            return bool(value)
        style = style.base_style
    return False


def _segments(paragraph):
    """[(text, bold, italic)] for every run, including runs inside hyperlinks."""
    style_bold = _style_flag(paragraph.style, "bold")
    style_italic = _style_flag(paragraph.style, "italic")
    segs = []
    for run in paragraph._p.iter(qn("w:r")):
        if any(a.tag == qn("w:txbxContent") for a in run.iterancestors()):
            continue
        text = ""
        for child in run:
            if child.tag == qn("w:t"):
                text += child.text or ""
            elif child.tag in (qn("w:tab"), qn("w:br")):
                text += " "
        if not text:
            continue
        rpr = run.find(qn("w:rPr"))
        bold, italic = _flag(rpr, "w:b"), _flag(rpr, "w:i")
        segs.append((
            text,
            style_bold if bold is None else bold,
            style_italic if italic is None else italic,
        ))
    return segs


def _clean(text):
    return re.sub(r"\s+", " ", text).strip()


class Para:
    def __init__(self, paragraph, index):
        self.index = index
        self.style = (paragraph.style.name or "").lower()
        self.segs = _segments(paragraph)
        raw = _clean("".join(s[0] for s in self.segs))
        pr = paragraph._p.find(qn("w:pPr"))
        has_num = pr is not None and pr.find(qn("w:numPr")) is not None
        starts_bullet = bool(raw) and raw[0] in "\u2022\u00b7\u25cf\u25aa"
        self.is_list = has_num or self.style.startswith("list") or starts_bullet
        self.text = re.sub(rf"^[{re.escape(BULLET_CHARS)}]\s+", "", raw) if self.is_list else raw

    @property
    def first_bold(self):
        for text, bold, _ in self.segs:
            if text.strip():
                return bold
        return False

    @property
    def all_italic(self):
        parts = [s for s in self.segs if s[0].strip()]
        return bool(parts) and all(s[2] for s in parts)


# ---------- classification ----------

def split_labeled(segs):
    """'**Backend:** Spring... **Frontend:** React' -> [(label, value), ...] or None."""
    items, label, value, state = [], [], [], None
    for text, bold, _ in segs:
        blank = not text.strip()
        if blank:
            (label if state == "label" else value if state == "value" else []).append(text)
            continue
        if bold:
            if state == "value":
                items.append((label, value))
                label, value = [], []
            elif state is None:
                label, value = [], []
            label.append(text)
            state = "label"
        else:
            if state is None:
                return None
            state = "value"
            value.append(text)
    if state is not None:
        items.append((label, value))
    result = []
    for lab, val in items:
        lab_text, val_text = _clean("".join(lab)), _clean("".join(val))
        if not lab_text.endswith(":") and val_text.startswith(":"):
            lab_text, val_text = lab_text + ":", _clean(val_text[1:])
        if not lab_text.endswith(":") or not val_text:
            return None
        result.append((lab_text[:-1].strip(), val_text))
    return result or None


def contact_type(text):
    low = text.lower()
    digits = re.sub(r"[\s\-()+]", "", text)
    if "@" in low or "email" in low:
        return "email"
    if "linkedin" in low:
        return "linkedin"
    if "github" in low:
        return "github"
    if "phone" in low or re.fullmatch(r"\d{8,15}", digits):
        return "phone"
    if low.startswith("http") or "www." in low:
        return "website"
    return "location"


def _fields(pairs):
    return [{"name": n, "text": t} for n, t in pairs]


def _parts_fields(text):
    return _fields((f"part{i + 1}", t) for i, t in enumerate(PART_SPLIT.split(text)))


def para_to_items(p):
    """One paragraph -> one item (a paragraph with several bold labels -> several).

    Every item remembers the paragraph it came from ("paras"), so the renderer
    can find the matching formatting in the template later.
    """
    items = _para_to_items(p)
    for it in items:
        it["paras"] = [p.index]
    return items


def _para_to_items(p):
    if p.is_list:
        return [{"role": "bullet", "fields": _fields([("text", p.text)])}]
    m = META_RE.match(p.text)
    if m:
        role = META_ROLE[m.group(1).lower()]
        return [{"role": role, "fields": _fields([("label", m.group(1)), ("value", m.group(2))])}]
    labeled = split_labeled(p.segs)
    if labeled:
        return [{"role": "labeled", "fields": _fields([("label", l), ("value", v)])}
                for l, v in labeled]
    if p.first_bold:
        return [{"role": "entry_title", "fields": _parts_fields(p.text)}]
    if p.all_italic:
        return [{"role": "subline", "fields": _parts_fields(p.text)}]
    return [{"role": "text", "fields": _fields([("text", p.text)])}]


def _merge_broken_text(items):
    """PDF-converted files split one paragraph in two; glue them back together."""
    merged = []
    for it in items:
        prev = merged[-1] if merged else None
        if (prev and it["role"] == "text" and prev["role"] == "text"
                and not prev["fields"][0]["text"].rstrip().endswith((".", "!", "?", ":", ";"))):
            prev["fields"][0]["text"] += " " + it["fields"][0]["text"]
            prev["paras"] = prev["paras"] + it["paras"]
        else:
            merged.append(it)
    return merged


def _is_heading(p):
    if p.is_list or not p.text or len(p.text) > 60:
        return False
    if canonical(p.text):
        return True
    if re.match(r"heading\s?[12]$", p.style):
        return True
    return p.text.isupper() and len(p.text.split()) <= 5 and p.first_bold


def _looks_like_contact(text):
    low = text.lower()
    return ("@" in text or "linkedin" in low or "github" in low
            or len(CONTACT_SPLIT.split(text)) >= 3
            or bool(re.search(r"\+?\d[\d\s\-()]{8,}\d", text)))


# ---------- whole document ----------

def parse_docx(path):
    """Return {'name','tagline','contact','sections':[...],'warnings':[...]}."""
    doc = Document(path)
    warnings = []
    paras = [Para(p, i) for i, p in enumerate(doc.paragraphs)]
    paras = [p for p in paras if p.text]
    if not paras:
        raise ValueError("The document has no readable text.")
    if doc.tables:
        warnings.append("This file contains tables. Text inside tables is not read yet.")

    result = {"name": None, "tagline": [], "contact": None, "sections": [], "warnings": warnings}
    result["name"] = {"role": "name", "fields": _fields([("text", paras[0].text)]),
                      "paras": [paras[0].index]}

    current = None
    for p in paras[1:]:
        if _is_heading(p):
            canon = canonical(p.text)
            if canon is None:
                warnings.append(f"Heading '{p.text}' is not in the synonym list, so it was kept as its own section.")
            current = {"title": p.text, "canonical": canon, "para": p.index, "paras": []}
            result["sections"].append(current)
        elif current is None:
            if _looks_like_contact(p.text) and result["contact"] is None:
                result["contact"] = {"role": "contact", "paras": [p.index], "fields": _fields(
                    (contact_type(part), part) for part in CONTACT_SPLIT.split(p.text))}
            else:
                result["tagline"].append({"role": "tagline", "paras": [p.index],
                                          "fields": _fields([("text", p.text)])})
        else:
            current["paras"].append(p)

    for sec in result["sections"]:
        items = [it for p in sec.pop("paras") for it in para_to_items(p)]
        items = _merge_broken_text(items)
        groups, cur = [], None
        for it in items:
            if it["role"] == "entry_title":
                cur = {"kind": "entry", "items": [it]}
                groups.append(cur)
            else:
                if cur is None:
                    cur = {"kind": "loose", "items": []}
                    groups.append(cur)
                cur["items"].append(it)
        sec["groups"] = groups

    if not result["sections"]:
        warnings.append("No section headings were found.")
    return result
