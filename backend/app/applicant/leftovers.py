"""Find placeholder text that was left behind in a finished CV."""
import re

from docx import Document

# {{W1_COMPANY}}   <<name>>   [COMPANY]   [Add your own education here]
PLACEHOLDER_RE = re.compile(r"\{\{[^}]*\}\}|<<[^>]*>>|\[[A-Za-z][^\]]{1,80}\]")


def _paragraph_texts(doc):
    for p in doc.paragraphs:
        yield p.text
    for section in doc.sections:
        for part in (section.header, section.footer):
            for p in part.paragraphs:
                yield p.text


def find_leftovers(docx_path):
    """Return the placeholders still present, e.g. ['{{W1_COMPANY}}']. Empty list = clean."""
    doc = Document(str(docx_path))
    found = []
    for text in _paragraph_texts(doc):
        for match in PLACEHOLDER_RE.findall(text):
            if match not in found:
                found.append(match)
    return found